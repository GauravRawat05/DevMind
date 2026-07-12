"""Report generation service — produces PDF and DOCX reports from analysis results.

Combines output from all four agents (Documentation, Code Review, Q&A,
Analytics) into a single professional downloadable report.
"""
from __future__ import annotations

import io
import json
import logging
import re
from datetime import datetime
from typing import Any

logger = logging.getLogger("devmind.report_generator")


# ---------------------------------------------------------------------------
# PDF Generation (using fpdf2)
# ---------------------------------------------------------------------------

def generate_pdf_report(
    job_id: str,
    repo_url: str,
    doc_output: str | None,
    review_output: list[dict[str, Any]] | None,
    analytics_output: dict[str, Any] | None,
    qa_history: list[dict[str, Any]] | None,
) -> bytes:
    """Generate a PDF report combining all agent outputs.

    Returns the PDF as bytes.
    """
    from fpdf import FPDF

    class DevMindPDF(FPDF):
        """Custom PDF class with header/footer branding."""

        def header(self):
            self.set_font("Helvetica", "B", 9)
            self.set_text_color(130, 130, 130)
            self.cell(0, 8, "DevMind AI Analysis Report", new_x="LMARGIN", new_y="NEXT", align="L")
            self.set_draw_color(0, 242, 254)
            self.set_line_width(0.4)
            self.line(10, self.get_y(), self.w - 10, self.get_y())
            self.ln(4)

        def footer(self):
            self.set_y(-15)
            self.set_font("Helvetica", "I", 8)
            self.set_text_color(160, 160, 160)
            self.cell(0, 10, f"Page {self.page_no()}/{{nb}}", align="C")

    pdf = DevMindPDF()
    pdf.alias_nb_pages()
    pdf.set_auto_page_break(auto=True, margin=20)

    # ---- Cover Page ----
    pdf.add_page()
    pdf.ln(40)
    pdf.set_font("Helvetica", "B", 28)
    pdf.set_text_color(0, 242, 254)
    pdf.cell(0, 14, "DevMind", new_x="LMARGIN", new_y="NEXT", align="C")
    pdf.set_font("Helvetica", "", 16)
    pdf.set_text_color(200, 200, 200)
    pdf.cell(0, 10, "AI Multi-Agent Analysis Report", new_x="LMARGIN", new_y="NEXT", align="C")
    pdf.ln(12)
    pdf.set_font("Helvetica", "", 11)
    pdf.set_text_color(160, 160, 160)
    pdf.cell(0, 8, _clean_for_pdf(f"Repository: {repo_url}"), new_x="LMARGIN", new_y="NEXT", align="C")
    pdf.cell(0, 8, _clean_for_pdf(f"Job ID: {job_id}"), new_x="LMARGIN", new_y="NEXT", align="C")
    pdf.cell(0, 8, _clean_for_pdf(f"Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}"), new_x="LMARGIN", new_y="NEXT", align="C")

    # ---- Helper to add section title ----
    def section_title(title: str):
        pdf.add_page()
        pdf.set_font("Helvetica", "B", 18)
        pdf.set_text_color(0, 242, 254)
        pdf.cell(0, 12, _clean_for_pdf(title), new_x="LMARGIN", new_y="NEXT")
        pdf.set_draw_color(0, 242, 254)
        pdf.set_line_width(0.3)
        pdf.line(10, pdf.get_y(), pdf.w - 10, pdf.get_y())
        pdf.ln(6)

    def body_text(text: str):
        pdf.set_font("Helvetica", "", 10)
        pdf.set_text_color(50, 50, 50)
        # Strip markdown formatting for PDF and clean unicode
        clean = _clean_for_pdf(_strip_markdown(text))
        pdf.multi_cell(0, 5.5, clean)
        pdf.ln(2)

    # ---- Section 1: Documentation ----
    section_title("1. Documentation (Doc Agent)")
    if doc_output:
        body_text(doc_output)
    else:
        body_text("No documentation was generated for this analysis.")

    # ---- Section 2: Code Review ----
    section_title("2. Code Review (Review Agent)")
    if review_output and len(review_output) > 0:
        # Summary
        severity_counts: dict[str, int] = {}
        for issue in review_output:
            sev = issue.get("severity", "info")
            severity_counts[sev] = severity_counts.get(sev, 0) + 1

        pdf.set_font("Helvetica", "B", 10)
        pdf.set_text_color(50, 50, 50)
        summary_parts = [f"{v} {k.upper()}" for k, v in severity_counts.items()]
        pdf.cell(0, 7, _clean_for_pdf(f"Total Issues Found: {len(review_output)}  ({', '.join(summary_parts)})"), new_x="LMARGIN", new_y="NEXT")
        pdf.ln(4)

        # Issues table
        col_widths = [50, 22, 55, 63]
        headers = ["File", "Severity", "Issue", "Suggestion"]

        pdf.set_font("Helvetica", "B", 9)
        pdf.set_fill_color(230, 230, 230)
        for i, h in enumerate(headers):
            pdf.cell(col_widths[i], 7, _clean_for_pdf(h), border=1, fill=True)
        pdf.ln()

        pdf.set_font("Helvetica", "", 8)
        for issue in review_output:
            sev = issue.get("severity", "info")
            if sev == "critical":
                pdf.set_text_color(220, 50, 50)
            elif sev == "warning":
                pdf.set_text_color(220, 160, 0)
            else:
                pdf.set_text_color(50, 50, 50)

            file_name = _clean_for_pdf(str(issue.get("file", "N/A"))[:30])
            message = _clean_for_pdf(str(issue.get("message", ""))[:50])
            suggestion = _clean_for_pdf(str(issue.get("suggestion", ""))[:55])

            row_data = [file_name, sev.upper(), message, suggestion]
            for i, val in enumerate(row_data):
                pdf.cell(col_widths[i], 6, _clean_for_pdf(val), border=1)
            pdf.ln()

        pdf.set_text_color(50, 50, 50)
    else:
        body_text("No code review issues were found, or the review agent did not produce results.")

    # ---- Section 3: Analytics ----
    section_title("3. Codebase Analytics (Analytics Agent)")
    if analytics_output:
        summary = analytics_output.get("summary", {})
        pdf.set_font("Helvetica", "B", 11)
        pdf.set_text_color(50, 50, 50)
        pdf.cell(0, 8, "Summary Metrics", new_x="LMARGIN", new_y="NEXT")
        pdf.ln(2)

        metrics = [
            ("Total Files", summary.get("total_files", "N/A")),
            ("Python Files", summary.get("total_python_files", "N/A")),
            ("Total Lines of Code", summary.get("total_loc", "N/A")),
            ("Average Complexity", f"{summary.get('avg_complexity', 'N/A')}"),
            ("Tech Debt Score", f"{analytics_output.get('tech_debt_score', 'N/A')}"),
        ]

        pdf.set_font("Helvetica", "", 10)
        for label, value in metrics:
            pdf.cell(70, 7, _clean_for_pdf(f"{label}:"), border=0)
            pdf.set_font("Helvetica", "B", 10)
            pdf.cell(0, 7, _clean_for_pdf(str(value)), new_x="LMARGIN", new_y="NEXT")
            pdf.set_font("Helvetica", "", 10)
        pdf.ln(4)

        # Top complex files
        top_files = analytics_output.get("top_complex_files", [])
        if top_files:
            pdf.set_font("Helvetica", "B", 11)
            pdf.cell(0, 8, "Top Complex Files", new_x="LMARGIN", new_y="NEXT")
            pdf.ln(2)

            pdf.set_font("Helvetica", "B", 9)
            pdf.set_fill_color(230, 230, 230)
            pdf.cell(90, 7, "File", border=1, fill=True)
            pdf.cell(30, 7, "LOC", border=1, fill=True)
            pdf.cell(35, 7, "Complexity", border=1, fill=True)
            pdf.ln()

            pdf.set_font("Helvetica", "", 8)
            pdf.set_text_color(50, 50, 50)
            for f in top_files[:15]:
                path = str(f.get("path", f.get("file", "unknown")))
                if len(path) > 50:
                    path = "..." + path[-47:]
                pdf.cell(90, 6, _clean_for_pdf(path), border=1)
                pdf.cell(30, 6, _clean_for_pdf(str(f.get("loc", f.get("lines", "-")))), border=1)
                pdf.cell(35, 6, _clean_for_pdf(str(f.get("complexity", "-"))), border=1)
                pdf.ln()
    else:
        body_text("No analytics data was generated for this analysis.")

    # ---- Section 4: Q&A History ----
    section_title("4. Q&A History (Q&A Agent)")
    if qa_history and len(qa_history) > 0:
        for idx, qa in enumerate(qa_history, 1):
            pdf.set_font("Helvetica", "B", 10)
            pdf.set_text_color(0, 100, 200)
            pdf.multi_cell(0, 6, _clean_for_pdf(f"Q{idx}: {qa.get('question', '')}"))
            pdf.ln(1)
            pdf.set_font("Helvetica", "", 9)
            pdf.set_text_color(50, 50, 50)
            answer = _clean_for_pdf(_strip_markdown(qa.get("answer", "")))
            pdf.multi_cell(0, 5, answer)
            pdf.ln(4)
    else:
        body_text("No Q&A interactions were recorded for this analysis.")

    # Output
    buf = io.BytesIO()
    pdf.output(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# DOCX Generation (using python-docx)
# ---------------------------------------------------------------------------

def generate_docx_report(
    job_id: str,
    repo_url: str,
    doc_output: str | None,
    review_output: list[dict[str, Any]] | None,
    analytics_output: dict[str, Any] | None,
    qa_history: list[dict[str, Any]] | None,
) -> bytes:
    """Generate a DOCX report combining all agent outputs.

    Returns the DOCX as bytes.
    """
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT

    doc = Document()

    # -- Document style defaults --
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(10)

    # ---- Cover ----
    doc.add_paragraph()
    doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("DevMind AI Analysis Report")
    run.font.size = Pt(26)
    run.font.color.rgb = RGBColor(0, 180, 220)
    run.bold = True

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(f"Repository: {repo_url}")
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(100, 100, 100)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run(f"Job ID: {job_id}  |  Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(140, 140, 140)

    doc.add_page_break()

    # ---- Section 1: Documentation ----
    h = doc.add_heading("1. Documentation (Doc Agent)", level=1)
    _color_heading(h, RGBColor(0, 180, 220))
    if doc_output:
        _add_markdown_to_docx(doc, doc_output)
    else:
        doc.add_paragraph("No documentation was generated for this analysis.")

    doc.add_page_break()

    # ---- Section 2: Code Review ----
    h = doc.add_heading("2. Code Review (Review Agent)", level=1)
    _color_heading(h, RGBColor(0, 180, 220))
    if review_output and len(review_output) > 0:
        # Summary
        severity_counts: dict[str, int] = {}
        for issue in review_output:
            sev = issue.get("severity", "info")
            severity_counts[sev] = severity_counts.get(sev, 0) + 1
        summary_parts = [f"{v} {k.upper()}" for k, v in severity_counts.items()]
        p = doc.add_paragraph()
        run = p.add_run(f"Total Issues Found: {len(review_output)}  ({', '.join(summary_parts)})")
        run.bold = True

        # Table
        table = doc.add_table(rows=1, cols=4)
        table.style = "Light Grid Accent 1"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        headers = ["File", "Severity", "Issue", "Suggestion"]
        for i, h_text in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = h_text
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True

        for issue in review_output:
            row = table.add_row().cells
            row[0].text = str(issue.get("file", "N/A"))
            sev = str(issue.get("severity", "info")).upper()
            row[1].text = sev
            row[2].text = str(issue.get("message", ""))
            row[3].text = str(issue.get("suggestion", ""))

            # Color the severity cell
            sev_color = {"CRITICAL": RGBColor(200, 40, 40), "WARNING": RGBColor(200, 150, 0)}
            if sev in sev_color:
                for paragraph in row[1].paragraphs:
                    for run in paragraph.runs:
                        run.font.color.rgb = sev_color[sev]
                        run.bold = True
    else:
        doc.add_paragraph("No code review issues were found, or the review agent did not produce results.")

    doc.add_page_break()

    # ---- Section 3: Analytics ----
    h = doc.add_heading("3. Codebase Analytics (Analytics Agent)", level=1)
    _color_heading(h, RGBColor(0, 180, 220))
    if analytics_output:
        summary = analytics_output.get("summary", {})

        doc.add_heading("Summary Metrics", level=2)
        metrics = [
            ("Total Files", summary.get("total_files", "N/A")),
            ("Python Files", summary.get("total_python_files", "N/A")),
            ("Total Lines of Code", summary.get("total_loc", "N/A")),
            ("Average Complexity", summary.get("avg_complexity", "N/A")),
            ("Tech Debt Score", analytics_output.get("tech_debt_score", "N/A")),
        ]
        for label, value in metrics:
            p = doc.add_paragraph()
            run = p.add_run(f"{label}: ")
            run.bold = True
            p.add_run(str(value))

        # Top complex files
        top_files = analytics_output.get("top_complex_files", [])
        if top_files:
            doc.add_heading("Top Complex Files", level=2)
            table = doc.add_table(rows=1, cols=3)
            table.style = "Light Grid Accent 1"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for i, h_text in enumerate(["File", "LOC", "Complexity"]):
                cell = table.rows[0].cells[i]
                cell.text = h_text
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True

            for f in top_files[:15]:
                row = table.add_row().cells
                row[0].text = str(f.get("path", f.get("file", "unknown")))
                row[1].text = str(f.get("loc", f.get("lines", "-")))
                row[2].text = str(f.get("complexity", "-"))
    else:
        doc.add_paragraph("No analytics data was generated for this analysis.")

    doc.add_page_break()

    # ---- Section 4: Q&A History ----
    h = doc.add_heading("4. Q&A History (Q&A Agent)", level=1)
    _color_heading(h, RGBColor(0, 180, 220))
    if qa_history and len(qa_history) > 0:
        for idx, qa in enumerate(qa_history, 1):
            p = doc.add_paragraph()
            run = p.add_run(f"Q{idx}: {qa.get('question', '')}")
            run.bold = True
            run.font.color.rgb = RGBColor(0, 80, 180)

            _add_markdown_to_docx(doc, qa.get("answer", ""))
            doc.add_paragraph()  # spacer
    else:
        doc.add_paragraph("No Q&A interactions were recorded for this analysis.")

    # Output
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Utility Helpers
# ---------------------------------------------------------------------------

def _clean_for_pdf(text: str) -> str:
    """Map unsupported Unicode characters to standard equivalents and decode safely to Latin-1."""
    if not text:
        return ""
    
    # Map common non-Latin1 unicode symbols to standard characters
    replacements = {
        "\u2022": "-",      # Bullet point
        "\u201c": '"',      # Left smart double quote
        "\u201d": '"',      # Right smart double quote
        "\u2018": "'",      # Left smart single quote
        "\u2019": "'",      # Right smart single quote
        "\u2013": "-",      # En dash
        "\u2014": "-",      # Em dash
        "\u2026": "...",    # Ellipsis
        "\xa0": " ",        # Non-breaking space
    }
    for orig, rep in replacements.items():
        text = text.replace(orig, rep)
    
    # Encode to Latin-1, replacing any remaining unsupported characters with ?
    try:
        return text.encode("latin-1", errors="replace").decode("latin-1")
    except Exception:
        return text.encode("ascii", errors="replace").decode("ascii")


def _strip_markdown(text: str) -> str:
    """Rough conversion of markdown to plain text for PDF rendering."""
    text = re.sub(r"```[\s\S]*?```", lambda m: m.group(0).strip("`").strip(), text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"#{1,6}\s+", "", text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    text = re.sub(r"\*([^*]+)\*", r"\1", text)
    text = re.sub(r"!\[[^\]]*\]\([^)]+\)", "", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    text = re.sub(r"^[-*+]\s+", "- ", text, flags=re.MULTILINE)
    text = re.sub(r"^\d+\.\s+", "- ", text, flags=re.MULTILINE)
    text = re.sub(r"^>\s+", "  ", text, flags=re.MULTILINE)
    text = re.sub(r"---+", "", text)
    return text.strip()


def _color_heading(heading, color: Any):
    """Apply a font color to all runs in a heading paragraph."""
    for run in heading.runs:
        run.font.color.rgb = color


def _add_markdown_to_docx(doc: Any, text: str):
    """Add markdown text to a DOCX document with basic formatting.

    Handles headings, bold, bullet points, and code blocks.
    """
    lines = text.split("\n")
    in_code_block = False
    code_buffer: list[str] = []

    for line in lines:
        # Code block toggle
        if line.strip().startswith("```"):
            if in_code_block:
                # End code block — flush buffer
                code_text = "\n".join(code_buffer)
                p = doc.add_paragraph()
                run = p.add_run(code_text)
                run.font.name = "Consolas"
                from docx.shared import Pt
                run.font.size = Pt(8)
                code_buffer = []
                in_code_block = False
            else:
                in_code_block = True
            continue

        if in_code_block:
            code_buffer.append(line)
            continue

        stripped = line.strip()
        if not stripped:
            continue

        # Headings
        heading_match = re.match(r"^(#{1,6})\s+(.*)", stripped)
        if heading_match:
            level = min(len(heading_match.group(1)), 4) + 1  # offset for sub-headings
            doc.add_heading(heading_match.group(2), level=level)
            continue

        # Bullet points
        bullet_match = re.match(r"^[-*+]\s+(.*)", stripped)
        if bullet_match:
            p = doc.add_paragraph(style="List Bullet")
            _add_inline_formatting(p, bullet_match.group(1))
            continue

        # Numbered list
        num_match = re.match(r"^\d+\.\s+(.*)", stripped)
        if num_match:
            p = doc.add_paragraph(style="List Number")
            _add_inline_formatting(p, num_match.group(1))
            continue

        # Regular paragraph
        p = doc.add_paragraph()
        _add_inline_formatting(p, stripped)


def _add_inline_formatting(paragraph: Any, text: str):
    """Add text with inline bold and code formatting to a paragraph."""
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            from docx.shared import Pt
            run.font.size = Pt(9)
        else:
            paragraph.add_run(part)
